import os
import sys
import pandas as pd
import customtkinter as ctk
from tkinter import filedialog, messagebox
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows

def process_data(raw_file_path):
    """
    Reads the raw data, filters for the last 10 hours, and aggregates KPIs per site.
    Returns:
        tuple: (list_of_10_hours, grouped_data_dict)
    """
    # Load raw data
    df = pd.read_excel(raw_file_path, sheet_name="Sheet0")

    # Extract unique hours and sort them (assuming format 'YYYY-MM-DD HH:MM:SS')
    unique_times = sorted(df['Begin Time'].unique())
    last_10_hours = unique_times[-10:] if len(unique_times) >= 10 else unique_times

    # Filter for the last 10 hours
    df_filtered = df[df['Begin Time'].isin(last_10_hours)]

    # Need to derive ORA_2G_TCH Availability(%) from other columns or default to 100 if it's missing in raw but present in template.
    # Actually, looking at raw columns, there is "ORA_2G_TCH Availability Normal TRXs" but no "ORA_2G_TCH Availability(%)"
    # Wait, sometimes it's missing, or maybe the template has it. Let's provide a fallback or calculate it.
    # In some datasets it's missing. We will use a fallback or try to find it.

    # KPIs to calculate (Output Name: (Raw Column Name, Aggregation function))
    kpis = {
        "Average of ORA_2G_TCH Availability(%)": ("ORA_2G_TCH Availability(%)", "mean"),
        "Average of ORA_2G_CSSR_CS_New(%)": ("ORA_2G_CSSR_CS_New(%)", "mean"),
        "Average of ORA_2G_Call_Drop_CS_New(%)": ("ORA_2G_Call_Drop_CS_New(%)", "mean"),
        "Sum of ORA_2G_CS_TRAFFIC": ("ORA_2G_CS_TRAFFIC", "sum"),
        "Average of ORA_2G_SDCCH_Blocking_New(%)": ("ORA_2G_SDCCH_Blocking_New(%)", "mean"),
        "Average of ORA_2G_TCH_Blocking_Rate_New(%)": ("ORA_2G_TCH_Blocking_Rate_New(%)", "mean")
    }

    # Check if 'ORA_2G_TCH Availability(%)' exists in raw data. If not, maybe use 'ORA_2G_TCH Availability Normal TRXs' or set to 100 as placeholder.
    if "ORA_2G_TCH Availability(%)" not in df_filtered.columns:
        if "ORA_2G_TCH Availability Normal TRXs" in df_filtered.columns:
            kpis["Average of ORA_2G_TCH Availability(%)"] = ("ORA_2G_TCH Availability Normal TRXs", "mean")
        else:
            df_filtered["ORA_2G_TCH Availability(%)"] = 100

    # Prepare a dictionary: {site_name: {kpi_name: {time: value}}}
    site_data = {}
    sites = df_filtered['SITE Name'].unique()

    for site in sites:
        df_site = df_filtered[df_filtered['SITE Name'] == site]
        site_data[site] = {}
        for kpi_out_name, (kpi_col, agg_func) in kpis.items():
            site_data[site][kpi_out_name] = {}
            for t in last_10_hours:
                df_t = df_site[df_site['Begin Time'] == t]
                if not df_t.empty:
                    val = df_t[kpi_col].mean() if agg_func == "mean" else df_t[kpi_col].sum()
                else:
                    val = None
                site_data[site][kpi_out_name][t] = val

    return last_10_hours, site_data

def apply_formatting(val, kpi_name):
    """
    Returns the appropriate hex color (string) based on the KPI value and rules.
    """
    if val is None or pd.isna(val):
        return "FFFFFF" # White for empty

    try:
        val = float(val)
    except ValueError:
        return "FFFFFF"

    # Green: 92D050 (approx), Red: FF0000, Faded Red: F08080 or similar.
    # From description:
    # Availability: >=98.5 Green, else faded red
    # CSSR CS: >= 98.5 Green, else faded red (described as < 98.5 rouge delave et vert au cas contraire)
    # Call drop CS: <= 0.7 Green, else red (> 0.7)
    # Traffic CS: always white
    # SDCCH / TCH BLOCKING: always white

    green = "00B050" # "92D050"
    faded_red = "FFC7CE" # Or "FF9999"
    red = "FF0000"
    white = "FFFFFF"

    if "Availability" in kpi_name:
        return green if val >= 98.5 else faded_red
    elif "CSSR" in kpi_name:
        return green if val >= 98.5 else faded_red
    elif "Call_Drop" in kpi_name:
        return green if val <= 0.7 else red
    else:
        return white

def export_to_excel(hours, site_data, output_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet2"

    # Write Header
    header = ["Row Labels"] + hours
    ws.append(header)

    # Header formatting
    header_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
    bold_font = Font(bold=True)
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

    for col_idx, _ in enumerate(header, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = bold_font
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center", vertical="center")

    current_row = 2
    for site, kpis in site_data.items():
        # Write Site Name
        ws.cell(row=current_row, column=1, value=site).font = bold_font
        ws.cell(row=current_row, column=1).fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
        ws.cell(row=current_row, column=1).border = thin_border
        for col_idx in range(2, len(hours) + 2):
            c = ws.cell(row=current_row, column=col_idx)
            c.fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
            c.border = thin_border
        current_row += 1

        # Write KPIs
        for kpi_name, times_vals in kpis.items():
            ws.cell(row=current_row, column=1, value=kpi_name).border = thin_border
            for col_idx, h in enumerate(hours, 2):
                val = times_vals.get(h, "")
                cell = ws.cell(row=current_row, column=col_idx, value=val)
                cell.border = thin_border
                # Apply conditional format
                color = apply_formatting(val, kpi_name)
                if color != "FFFFFF":
                    cell.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
                if isinstance(val, (int, float)):
                    cell.number_format = '0.00'
            current_row += 1

    # Adjust column widths
    ws.column_dimensions['A'].width = 45
    for col_idx in range(2, len(hours) + 2):
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = 15

    wb.save(output_path)

# Main Application Class
class App(ctk.CTk):
    def __init__(self):
        super().__init__()
        self.title("2G KPI Report Generator")
        self.geometry("650x450")

        # Appearance
        ctk.set_appearance_mode("System")
        ctk.set_default_color_theme("blue")

        self.raw_file_path = ""
        self.output_dir_path = ""

        self.build_ui()

    def build_ui(self):
        # Title Label
        title = ctk.CTkLabel(self, text="2G KPI Report Generator", font=ctk.CTkFont(size=24, weight="bold"))
        title.pack(pady=20)

        # Frame for Raw File
        frame1 = ctk.CTkFrame(self, fg_color="transparent")
        frame1.pack(pady=10, padx=20, fill="x")

        self.lbl_raw = ctk.CTkLabel(frame1, text="No raw data file selected.", anchor="w")
        self.lbl_raw.pack(side="left", padx=10, fill="x", expand=True)

        btn_raw = ctk.CTkButton(frame1, text="Select Raw Data", command=self.select_raw_file)
        btn_raw.pack(side="right", padx=10)

        # Frame for Output Dir
        frame2 = ctk.CTkFrame(self, fg_color="transparent")
        frame2.pack(pady=10, padx=20, fill="x")

        self.lbl_out = ctk.CTkLabel(frame2, text="No output directory selected.", anchor="w")
        self.lbl_out.pack(side="left", padx=10, fill="x", expand=True)

        btn_out = ctk.CTkButton(frame2, text="Select Output Folder", command=self.select_output_dir)
        btn_out.pack(side="right", padx=10)

        # Frame for Output Filename
        frame3 = ctk.CTkFrame(self, fg_color="transparent")
        frame3.pack(pady=10, padx=20, fill="x")

        lbl_name = ctk.CTkLabel(frame3, text="Output Filename (without .xlsx):")
        lbl_name.pack(side="left", padx=10)

        self.entry_name = ctk.CTkEntry(frame3, width=200)
        self.entry_name.insert(0, "KPI_Report_Output")
        self.entry_name.pack(side="left", padx=10)

        # Generate Button
        self.btn_generate = ctk.CTkButton(self, text="Generate Excel Report", command=self.generate_report, height=40, font=ctk.CTkFont(size=16, weight="bold"))
        self.btn_generate.pack(pady=30)

        # Status Label
        self.lbl_status = ctk.CTkLabel(self, text="", text_color="gray")
        self.lbl_status.pack(pady=10)

    def select_raw_file(self):
        filetypes = [("Excel files", "*.xlsx *.xls")]
        filepath = filedialog.askopenfilename(title="Select Raw Data File", filetypes=filetypes)
        if filepath:
            self.raw_file_path = filepath
            self.lbl_raw.configure(text=f"Selected: {os.path.basename(filepath)}")

    def select_output_dir(self):
        dirpath = filedialog.askdirectory(title="Select Output Folder")
        if dirpath:
            self.output_dir_path = dirpath
            self.lbl_out.configure(text=f"Folder: {dirpath}")

    def generate_report(self):
        if not self.raw_file_path:
            messagebox.showwarning("Warning", "Please select a raw data file.")
            return
        if not self.output_dir_path:
            messagebox.showwarning("Warning", "Please select an output folder.")
            return

        filename = self.entry_name.get().strip()
        if not filename:
            messagebox.showwarning("Warning", "Please enter an output filename.")
            return

        if not filename.endswith(".xlsx"):
            filename += ".xlsx"

        output_path = os.path.join(self.output_dir_path, filename)

        self.lbl_status.configure(text="Processing data...", text_color="blue")
        self.update()

        try:
            hours, site_data = process_data(self.raw_file_path)
            export_to_excel(hours, site_data, output_path)

            self.lbl_status.configure(text=f"Success! Saved to {filename}", text_color="green")
            messagebox.showinfo("Success", f"Report successfully generated at:\n{output_path}")
        except Exception as e:
            self.lbl_status.configure(text="Error occurred.", text_color="red")
            messagebox.showerror("Error", f"An error occurred during generation:\n{str(e)}")

if __name__ == "__main__":
    app = App()
    app.mainloop()
