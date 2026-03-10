import pandas as pd
from datetime import datetime
import openpyxl
from openpyxl.styles import PatternFill, Border, Side, Alignment, Font
from openpyxl.utils import get_column_letter
import os

def determine_site_column(df):
    possible_columns = ['Group', 'SITE Name', 'Cell', 'Site', 'NE Name']
    for col in possible_columns:
        if col in df.columns:
            return col
    return None

def process_raw_data(file_path):
    print(f"Loading data from {file_path}")
    xls = pd.ExcelFile(file_path)
    df = pd.read_excel(xls, sheet_name=xls.sheet_names[0])

    site_col = determine_site_column(df)
    if not site_col:
        raise ValueError("Could not determine site or city column in the data.")

    df['Begin Time'] = pd.to_datetime(df['Begin Time'])

    # Filter for the last 10 hours
    unique_times = sorted(df['Begin Time'].unique())
    last_10 = unique_times[-10:]
    df_filtered = df[df['Begin Time'].isin(last_10)].copy()

    if 'ORA_2G_TCH Availability Normal TRXs' in df_filtered.columns:
        if df_filtered['ORA_2G_TCH Availability Normal TRXs'].max() <= 1.0:
            df_filtered['ORA_2G_TCH Availability Normal TRXs'] = df_filtered['ORA_2G_TCH Availability Normal TRXs'] * 100

    percentage_cols = [
        'ORA_2G_CSSR_CS_New(%)',
        'ORA_2G_Call_Drop_CS_New(%)',
        'ORA_2G_SDCCH_Blocking_New(%)',
        'ORA_2G_TCH_Blocking_Rate_New(%)'
    ]
    for col in percentage_cols:
        if col in df_filtered.columns:
            # Let's check max value to be safe, though some % columns might actually exceed 1 if they are already %
            # If max > 1, it might already be in %. If not, we multiply.
            # However, CSSR and Drop Rate in the raw data might sometimes have a value > 1 (e.g. 99.61)
            # So if max is <= 1.0, it's a fraction.
            if df_filtered[col].max() <= 1.0:
                df_filtered[col] = df_filtered[col] * 100

    # Sort data
    df_filtered = df_filtered.sort_values(by=[site_col, 'Begin Time'])
    return df_filtered, site_col, last_10

def generate_report(file_path, event_name):
    try:
        df, site_col, last_10 = process_raw_data(file_path)
    except Exception as e:
        return f"Error processing file: {e}"

    # Generate output file name
    date_str = datetime.now().strftime("%d%m%Y")
    hour_str = datetime.now().strftime("%H%M")
    output_filename = f"KEA_2G_{event_name}_{date_str}_{hour_str}.xlsx"

    # Usually we save in the user's Downloads directory
    if os.name == 'nt': # Windows
        downloads_dir = os.path.join(os.environ['USERPROFILE'], 'Downloads')
    else: # macOS / Linux
        downloads_dir = os.path.join(os.path.expanduser('~'), 'Downloads')

    # create downloads directory if it doesn't exist
    if not os.path.exists(downloads_dir):
        os.makedirs(downloads_dir)

    output_filepath = os.path.join(downloads_dir, output_filename)

    # Create workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    # Define styles matching template exactly
    green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid") # Template Green (FFC6EFCE)
    pink_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid") # Template Rose/Pink (FFFFC7CE)
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color='000000'),
        right=Side(style='thin', color='000000'),
        top=Side(style='thin', color='000000'),
        bottom=Side(style='thin', color='000000')
    )

    bold_font = Font(bold=True)
    center_aligned = Alignment(horizontal='center', vertical='center')

    # Write header
    ws.cell(row=1, column=1, value="Row Labels").font = bold_font
    ws.cell(row=1, column=1).border = thin_border

    for i, t in enumerate(last_10):
        c = ws.cell(row=1, column=2+i, value=t.strftime("%Y-%m-%d %H:%M:%S"))
        c.font = bold_font
        c.border = thin_border
        c.alignment = center_aligned

    # Group data by site
    current_row = 2
    for site, group in df.groupby(site_col):
        # Write site name
        site_cell = ws.cell(row=current_row, column=1, value=site)
        site_cell.font = bold_font
        for i in range(1, len(last_10) + 2):
            ws.cell(row=current_row, column=i).border = thin_border
            if i > 1:
                ws.cell(row=current_row, column=i, value="")
        current_row += 1

        # Prepare KPI data rows
        # kpis define how to format colors.
        # lambda val: fill based on condition
        def get_kpi_configs():
            return [
                {"label": "Average of ORA_2G_TCH Availability(%)", "col": "ORA_2G_TCH Availability Normal TRXs", "cond": lambda v: green_fill if v >= 98.5 else pink_fill},
                {"label": "Average of ORA_2G_CSSR_CS_New(%)", "col": "ORA_2G_CSSR_CS_New(%)", "cond": lambda v: green_fill if v >= 98.5 else pink_fill},
                {"label": "Average of ORA_2G_Call_Drop_CS_New(%)", "col": "ORA_2G_Call_Drop_CS_New(%)", "cond": lambda v: green_fill if v <= 0.7 else pink_fill},
                {"label": "Sum of ORA_2G_CS_TRAFFIC", "col": "ORA_2G_CS_TRAFFIC", "cond": lambda v: white_fill},
                {"label": "Average of ORA_2G_SDCCH_Blocking_New(%)", "col": "ORA_2G_SDCCH_Blocking_New(%)", "cond": lambda v: white_fill},
                {"label": "Average of ORA_2G_TCH_Blocking_Rate_New(%)", "col": "ORA_2G_TCH_Blocking_Rate_New(%)", "cond": lambda v: white_fill}
            ]

        for kpi in get_kpi_configs():
            # Write KPI label
            c_label = ws.cell(row=current_row, column=1, value=kpi["label"])
            c_label.border = thin_border

            # Write values
            for i, t in enumerate(last_10):
                # find value for this time
                val_row = group[group['Begin Time'] == t]
                val = ""
                cell = ws.cell(row=current_row, column=2+i)
                if not val_row.empty:
                    if kpi["col"] in val_row.columns:
                        raw_val = val_row[kpi["col"]].values[0]
                        if pd.notna(raw_val):
                            val = round(float(raw_val), 2)

                if val != "":
                    cell.value = val
                    cell.number_format = '0.00'
                    cell.fill = kpi["cond"](val)
                else:
                    cell.value = ""
                    cell.fill = white_fill

                cell.border = thin_border
                cell.alignment = center_aligned

            current_row += 1

        # Add empty row to separate site blocks, keeping aesthetics as in the template
        current_row += 1

    # Adjust column widths
    ws.column_dimensions['A'].width = 45
    for i in range(2, len(last_10) + 2):
        col_letter = get_column_letter(i)
        ws.column_dimensions[col_letter].width = 20

    wb.save(output_filepath)
    return f"Report successfully generated at:\n{output_filepath}"

if __name__ == "__main__":
    # Test
    res = generate_report("Performance Management-History Query-2G_KPI_Reporting_Template-DFBG6870-20260309081217.xlsx", "TestEvent")
    print(res)
